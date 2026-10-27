import asyncio
import hashlib
import os
import shutil
from pathlib import Path
from datetime import datetime
from Core.clock import now
from Core.results import result,save,saved
from .son_function.nickname_segmenter import split_nickname
from .son_function.align_slots import align_slots
from .son_function.write_in_excel import write_event
SLOTS=("元谥","亨谥","利谥","贞谥","常谥","平谥")
def remarks(changes):
    parts=[]
    for c in changes:
        name=SLOTS[c["slot"]]
        if c["kind"]=="delete":parts.append(f'{name}“{c["old"]}”删除')
        elif c["kind"]=="insert":parts.append(f'{name}新增“{c["new"]}”')
        else:parts.append(f'{name}“{c["old"]}”→“{c["new"]}”')
    return "；".join(parts)
def snapshot(source,dest):
    dest.parent.mkdir(parents=True,exist_ok=True)
    if not dest.exists():
        temp=dest.with_suffix(".partial");shutil.copyfile(source,temp);os.replace(temp,dest)
    return str(dest)
async def run(task,ctx):
    previous=await saved(task,ctx)
    if previous is not None:return previous
    gid=ctx.db.check_group(task["group_id"]);uid=task["user_id"];p=task["params"]
    new=p["new_nickname"];observed=datetime.fromisoformat(p["observed_at"])
    rows=await ctx.db.read("log",group_id=gid,filters={"task_id":task["task_id"]},limit=1)
    state=rows[0]["result"]
    members=await ctx.db.read("members",group_id=gid,filters={"user_id":uid},limit=1)
    member=members[0] if members else None
    earlier=await ctx.db.read("log",group_id=gid,filters={"request_type":"memento","user_id":uid},order_by=(("created_at","desc"),))
    blocked=next((r for r in earlier if str(r["task_id"])!=str(task["task_id"]) and r["result"].get("alignment_required")),None)
    if blocked and not p.get("manual_slots"):
        return result("error",error_code="nickname_alignment_required")
    if not state.get("event"):
        # Work from the last successfully exported state, never infer slots from an unprocessed observation.
        completed=next((r for r in earlier if r["business_status"]=="completed" and r["result"].get("event")),None)
        old_event=completed["result"]["event"] if completed else None
        if old_event and observed<datetime.fromisoformat(old_event["observed_at"]):
            return await save(task,ctx,result("skipped",result_id=task["task_id"]))
        old=old_event["new_nickname"] if old_event else p.get("old_nickname")
        initial=p.get("is_initial",False) and old_event is None
        old_slots=old_event["new_slots"] if old_event else (member or {}).get("slots",[None]*6)
        if old==new and old_event and not p.get("manual_slots"):
            return await save(task,ctx,result("skipped",result_id=task["task_id"]))
        try:
            if p.get("manual_slots") is not None:
                slots=p["manual_slots"]
                if not isinstance(slots,list) or len(slots)!=6 or any(x is not None and (not isinstance(x,str) or not x) for x in slots) or "".join(x or "" for x in slots)!=new:
                    raise ValueError("manual slots invalid")
                aligned={"status":"aligned","slots":slots,"changes":[]}
            else:aligned=align_slots([None]*6 if initial else old_slots,split_nickname(new,[None]*6 if initial else old_slots))
            if aligned["status"]!="aligned":raise ValueError("ambiguous")
        except ValueError:
            await ctx.db.write("log",[{"task_id":task["task_id"],"result":{"alignment_required":True,"new_nickname":new},"step":"alignment_required"}],group_id=gid,mode="update",key_fields=("task_id",))
            return result("error",error_code="nickname_alignment_required")
        sheet=(member or {}).get("sheet_name") or uid
        workbook=ctx.settings.data_root/"groups"/gid/"memento"/"鸿爪录.xlsx"
        if not (member or {}).get("sheet_name") and workbook.exists():
            from openpyxl import load_workbook
            book=await asyncio.to_thread(load_workbook,workbook,read_only=True)
            conflict=sheet in book.sheetnames;book.close()
            if conflict: return result("error",error_code="member_sheet_binding_required")
        event={"event_id":str(task["task_id"]),"user_id":uid,"new_nickname":new,"old_nickname":None if initial else old,
               "new_slots":aligned["slots"],"observed_at":observed.isoformat(),"remark":"初创" if initial else remarks(aligned["changes"]),
               "display_name":(member or {}).get("display_name") or new[:4]}
        state={"event":event,"sheet_name":sheet,"initial":initial,"phase":"aligned"}
        await ctx.db.write("log",[{"task_id":task["task_id"],"result":state,"step":"aligned"}],group_id=gid,mode="update",key_fields=("task_id",))
    event=state["event"];sheet=state["sheet_name"]
    workbook=ctx.settings.data_root/"groups"/gid/"memento"/"鸿爪录.xlsx"
    await asyncio.to_thread(write_event,workbook,sheet,event)
    actions=[]
    if not state["initial"]:
        dest=ctx.settings.data_root/"groups"/gid/"outgoing"/str(task["task_id"])/"0"/"鸿爪录.xlsx"
        path=await asyncio.to_thread(snapshot,workbook,dest)
        actions=[{"type":"upload_file","path":path,"name":"鸿爪录.xlsx"},{"type":"send_message","text":f'鸿爪已录：{event["display_name"]}\n旧名：{event["old_nickname"] or ""}\n新名：{new}\n《鸿爪录》已更新。'}]
    value=result(actions=actions,result_id=task["task_id"],event=event,sheet_name=sheet,phase="exported")
    async with ctx.db.transaction(gid) as tx:
        await ctx.db.write("members",[{"user_id":uid,"nickname":new,"slots":event["new_slots"],"sheet_name":sheet,"updated_at":now()}],group_id=gid,mode="upsert",key_fields=("user_id",),transaction=tx)
        if p.get("manual_slots"):
            for record in earlier:
                if record["result"].get("alignment_required"):
                    r={**record["result"],"alignment_required":False,"resolved_by":str(task["task_id"])}
                    await ctx.db.write("log",[{"task_id":record["task_id"],"result":r}],group_id=gid,mode="update",key_fields=("task_id",),transaction=tx)
        await save(task,ctx,value,tx=tx)
    return value
