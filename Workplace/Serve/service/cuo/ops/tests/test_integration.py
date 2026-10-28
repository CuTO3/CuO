import asyncio,copy,json
from datetime import datetime,timedelta
from uuid import uuid4
import pytest
from psycopg import sql
from Core.clock import BEIJING,business_date,now
from Core.results import result,save
from Integration.napcat.delivery import deliver,reconcile,RetryDelivery
from Integration.napcat.client import DeliveryUnknown,NotConnected
from Core.task_queue import TaskQueue
from Core.cache import CacheManager
from Workflow._receiver_ import receive_event
from Workflow.main import dispatch
async def task(ctx,kind="divine",**params):
    gid=ctx.gid;t=datetime(2026,10,26,2,tzinfo=BEIJING)
    member=await ctx.db.read("members",group_id=gid,filters={"user_id":"123"})
    if not member:await ctx.db.write("members",[{"user_id":"123","nickname":"测试群友"}],group_id=gid)
    row=(await ctx.db.write("log",[{"task_id":str(uuid4()),"triggered_at":t,"business_date":business_date(t),"request_type":kind,"user_id":"123","original_text":"你好","params":{"snapshot_seq":1000,**params}}],group_id=gid))[0]
    return {**row,"group_id":gid}
def event(ctx,mid=1,text="你好",at=True):
    return {"post_type":"message","message_type":"group","group_id":ctx.gid,"self_id":"999","user_id":"123","message_id":mid,"time":1792951200,"sender":{"nickname":"测试群友","user_id":"123"},"message":([{"type":"at","data":{"qq":"999"}}] if at else [])+[{"type":"text","data":{"text":text}}]}
async def stored(ctx,t):
    return (await ctx.db.read("log",group_id=ctx.gid,filters={"task_id":t["task_id"]}))[0]
async def test_receiver_dedup_and_order(ctx):
    e=event(ctx);e["message"].append({"type":"image","data":{"url":"https://example.invalid/a.png"}})
    a=await receive_event(e,now(),context=ctx);b=await receive_event(e,now(),context=ctx)
    assert a["status"]=="accepted" and b["status"]=="duplicate"
    logs=await ctx.db.read("log",group_id=ctx.gid,filters={})
    assert len(logs)==3
    assert sorted((l["params"]["queue"]["order"],l["request_type"]) for l in logs)==[(0,"memento"),(1,"media_describe"),(2,"message")]
    assert not ctx.model.calls
async def test_plain_at_and_self(ctx):
    await receive_event(event(ctx,text="@999",at=False),now(),context=ctx)
    e=event(ctx,mid=2);e["user_id"]="999";e["sender"]["user_id"]="999";e["post_type"]="message_sent"
    await receive_event(e,now(),context=ctx)
    assert not await ctx.db.read("log",group_id=ctx.gid,filters={"request_type":"message"})
    rows=await ctx.db.read("messages",group_id=ctx.gid,filters={"message_id":"2"})
    assert rows[0]["generated_by"]=="bot"
async def test_group_isolation_and_validation(ctx):
    await task(ctx)
    assert await ctx.db.read("members",group_id=ctx.other,filters={})==[]
    with pytest.raises(ValueError):await ctx.db.read("members",group_id="evil",filters={})
    with pytest.raises((TypeError,ValueError)):await ctx.db.write("members",[{"user_id":"456","oops":"bad"}],group_id=ctx.gid)
    assert not await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"456"})
async def test_rollback(ctx):
    with pytest.raises(RuntimeError):
        async with ctx.db.transaction(ctx.gid) as tx:
            await ctx.db.write("members",[{"user_id":"55","nickname":"rollback"}],group_id=ctx.gid,transaction=tx)
            raise RuntimeError()
    assert not await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"55"})
async def test_divine_once_and_recovery(ctx,monkeypatch):
    monkeypatch.setattr("Function.divine.divine.draw_dictum_luck",lambda dictum:0)
    from Function.divine.divine import run
    t=await task(ctx)
    first=await run(t,ctx);again=await run(t,ctx)
    assert first==again
    record=(await ctx.db.read("divinations",group_id=ctx.gid,filters={}))[0]
    m=(await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"123"}))[0]
    assert m["fortune"]==record["fortune_after"]==record["luck_delta"]
    second=await run(await task(ctx),ctx)
    assert second["error_code"]=="insufficient_fortune"
    assert second["actions"][0]["text"]=="今日缘法已尽，且待明日再问。"
    assert second["actions"][1]["text"]==record["content"]
    assert not ctx.model.calls
    ctx.onebot.fail=NotConnected()
    with pytest.raises(RetryDelivery):await deliver(t,ctx,first)
    row=await stored(ctx,t);ctx.onebot.fail=None
    await dispatch({**row,"group_id":ctx.gid},ctx)
    m2=(await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"123"}))[0]
    assert m2["fortune"]==m["fortune"] and (await stored(ctx,t))["status"]=="completed"
async def test_unknown_not_resent_and_late_receipt(ctx):
    t=await task(ctx,"conversation");value=await save(t,ctx,result(actions=[{"type":"send_message","text":"x"}]))
    ctx.onebot.fail=DeliveryUnknown()
    await deliver(t,ctx,value)
    row=await stored(ctx,t);assert row["status"]=="needs_review"
    await dispatch({**row,"group_id":ctx.gid},ctx)
    assert len(ctx.onebot.calls)==1
    echo=row["result"]["actions"][0]["echo"]
    await reconcile(ctx,{"echo":echo,"status":"ok","retcode":0,"data":{"message_id":"567"}})
    assert (await stored(ctx,t))["status"]=="completed"
async def test_crash_sending_recovery(ctx):
    t=await task(ctx,"conversation")
    await ctx.db.write("log",[{"task_id":t["task_id"],"status":"running","business_status":"completed","result":{"actions":[{"status":"sending","action_id":"x"}]}}],group_id=ctx.gid,mode="update",key_fields=("task_id",))
    queue=TaskQueue(ctx,dispatch);await queue.recover()
    row=await stored(ctx,t)
    assert row["status"]=="needs_review" and row["result"]["actions"][0]["status"]=="unknown"
async def test_worker_exclusion(ctx):
    q1=TaskQueue(ctx,dispatch);q2=TaskQueue(ctx,dispatch)
    try:
        await q1.start(launch=False)
        with pytest.raises(RuntimeError):await q2.start(launch=False)
    finally:await q1.stop();await q2.stop()
async def test_conversation_saved_once(ctx):
    from Function.conversation.conversation import run
    t=await task(ctx,"conversation");a=await run(t,ctx);b=await run(t,ctx)
    assert a==b and len(ctx.model.calls)==1 and a["status"]=="ok"
async def test_memento_idempotent_workbook(ctx):
    from Function.memento.memento import run
    from openpyxl import load_workbook
    t=await task(ctx,"memento",new_nickname="马铃薯",old_nickname=None,is_initial=True,observed_at=now().isoformat())
    a=await run(t,ctx);assert a["actions"]==[]
    await run(t,ctx)
    path=ctx.settings.data_root/"groups"/ctx.gid/"memento/鸿爪录.xlsx"
    wb=load_workbook(path);assert wb["123"].max_row==2 and wb["_cuo_export"].max_row==2;wb.close()
async def test_cache_commit_and_rollback(ctx):
    cache=CacheManager(ctx.settings,ctx.db);ctx.cache=cache;ctx.db.set_cache(cache)
    await receive_event(event(ctx),now(),context=ctx)
    messages=await ctx.db.read("messages",group_id=ctx.gid,filters={});day=messages[0]["business_date"]
    assert len(await cache.read_messages(ctx.gid,day))==1
    with pytest.raises(RuntimeError):
        async with ctx.db.transaction(ctx.gid) as tx:
            await ctx.db.write("messages",[{"message_id":"bad","sent_at":now(),"business_date":day,"user_id":"123"}],group_id=ctx.gid,transaction=tx)
            raise RuntimeError()
    assert len(await cache.read_messages(ctx.gid,day))==1
    cache.paths(ctx.gid,day)["messages"].write_text("broken",encoding="utf-8")
    assert len(await cache.read_messages(ctx.gid,day))==1
async def test_summary_evidence_and_reuse(ctx):
    from Function.summary.summary import run
    t=await task(ctx,"summary",request_text="总结一下")
    await ctx.db.write("messages",[{"message_id":"41","sent_at":t["triggered_at"]-timedelta(hours=1),"business_date":t["business_date"],"user_id":"123","nickname":"测试群友","text":"今晚吃火锅","segments":[]}],group_id=ctx.gid)
    ctx.model.reply=json.dumps({"events":[{"title":"晚饭","body":"讨论吃火锅","refs":["m00001"],"quotes":[{"ref":"m00001","text":"吃火锅"}]}]},ensure_ascii=False)
    a=await run(t,ctx);assert a["status"]=="ok"
    b=await run(await task(ctx,"summary",request_text="总结一下"),ctx)
    assert b["result_id"]==a["result_id"] and b["reused"] and len(ctx.model.calls)==1
async def test_admin_stale_version(ctx):
    from Integration.console.admin import execute,version
    t=await task(ctx,"admin_update",operation="display_name",user_id="123",expected_version="stale",display_name="新名")
    r=await execute(t,ctx);assert r["error_code"]=="version_conflict"
    member=(await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"123"}))[0]
    t=await task(ctx,"admin_update",operation="display_name",user_id="123",expected_version=version(member),display_name="新名")
    assert (await execute(t,ctx))["status"]=="ok"

@pytest.mark.parametrize('value',[0,-12,42,9007199254740991])
async def test_admin_fortune_persists_with_audit(ctx,value):
    from Integration.console.admin import execute,version
    await task(ctx)
    member=(await ctx.db.read('members',group_id=ctx.gid,filters={'user_id':'123'}))[0]
    t=await task(ctx,'admin_update',operation='fortune',user_id='123',fortune=value,expected_version=version(member))
    r=await execute(t,ctx)
    updated=(await ctx.db.read('members',group_id=ctx.gid,filters={'user_id':'123'}))[0]
    assert updated['fortune']==value and updated['nickname']==member['nickname']
    assert r['status']=='ok' and r['actions']==[] and r['fortune_before']==member['fortune'] and r['fortune_after']==value
    assert (await stored(ctx,t))['result']==r
    assert await ctx.db.read('members',group_id=ctx.other,filters={'user_id':'123'})==[]
    assert not ctx.onebot.calls

async def test_admin_fortune_concurrent_change_is_not_overwritten(ctx):
    from Integration.console.admin import execute,version
    await task(ctx)
    member=(await ctx.db.read('members',group_id=ctx.gid,filters={'user_id':'123'}))[0]
    t=await task(ctx,'admin_update',operation='fortune',user_id='123',fortune=99,expected_version=version(member))
    await ctx.db.write('members',[{'user_id':'123','fortune':7}],group_id=ctx.gid,mode='update',key_fields=('user_id',))
    r=await execute(t,ctx)
    assert r['error_code']=='version_conflict'
    assert (await ctx.db.read('members',group_id=ctx.gid,filters={'user_id':'123'}))[0]['fortune']==7

async def test_admin_fortune_missing_member_is_not_created(ctx):
    from Integration.console.admin import execute
    t=await task(ctx,'admin_update',operation='fortune',user_id='456',fortune=99,expected_version='missing')
    assert (await execute(t,ctx))['error_code']=='version_conflict'
    assert await ctx.db.read('members',group_id=ctx.gid,filters={'user_id':'456'})==[]
async def test_archive_no_send_and_late_revision(ctx):
    from Core.archive import schedule,run
    # Use yesterday relative to the actual scheduler clock, not the test trigger.
    d=business_date(now())-timedelta(days=1)
    await ctx.db.write("messages",[{"message_id":"71","sent_at":now()-timedelta(days=1),"business_date":d,"user_id":"123","text":"晚饭"}],group_id=ctx.gid)
    await schedule(ctx)
    rows=await ctx.db.read("log",group_id=ctx.gid,filters={"request_type":"archive_day"})
    assert len(rows)==1
    r=await run({**rows[0],"group_id":ctx.gid},ctx);assert r["actions"]==[] and r["message_count"]==1
    await ctx.db.write("messages",[{"message_id":"72","sent_at":now()-timedelta(days=1),"business_date":d,"user_id":"123","text":"火锅"}],group_id=ctx.gid)
    await schedule(ctx)
    rows=await ctx.db.read("log",group_id=ctx.gid,filters={"request_type":"archive_day"})
    assert len(rows)==2 and any(r["result"].get("source_dirty") for r in rows)

async def test_divine_repeat_fee_latest_and_idempotency(ctx,monkeypatch):
    from Function.divine.divine import run
    from ops.database.scripts.migrate import migrate
    monkeypatch.setattr("Function.divine.divine.draw_dictum_luck",lambda dictum:0)
    await run(await task(ctx),ctx)
    await ctx.db.write("members",[{"user_id":"123","fortune":4}],group_id=ctx.gid,mode="update",key_fields=("user_id",))
    for expected in (2,0):
        t=await task(ctx)
        value=await run(t,ctx)
        assert value["status"]=="ok"
        assert value["actions"][0]["text"].startswith("往日福缘深厚，尚可再问天机。")
        assert await run(t,ctx)==value
        member=(await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"123"}))[0]
        assert member["fortune"]==expected
    latest=value
    denied=await run(await task(ctx),ctx)
    assert denied["error_code"]=="insufficient_fortune"
    assert denied["result_id"]==latest["result_id"]
    assert denied["actions"][1]==latest["actions"][0]
    assert len(await ctx.db.read("divinations",group_id=ctx.gid,filters={}))==3
    await migrate(ctx.db)

async def test_divine_repeat_resume_does_not_charge_twice(ctx,monkeypatch):
    from Function.divine import divine
    monkeypatch.setattr(divine,"draw_dictum_luck",lambda dictum:0)
    await divine.run(await task(ctx),ctx)
    await ctx.db.write("members",[{"user_id":"123","fortune":2}],group_id=ctx.gid,mode="update",key_fields=("user_id",))
    t=await task(ctx)
    original=divine.generate
    async def fail(*args):raise RuntimeError("temporary model failure")
    monkeypatch.setattr(divine,"generate",fail)
    with pytest.raises(RuntimeError):await divine.run(t,ctx)
    pending=await divine.run(await task(ctx),ctx)
    assert pending["error_code"]=="divination_pending"
    monkeypatch.setattr(divine,"generate",original)
    assert (await divine.run(t,ctx))["status"]=="ok"
    member=(await ctx.db.read("members",group_id=ctx.gid,filters={"user_id":"123"}))[0]
    assert member["fortune"]==0

async def test_divine_dictum_direction_and_repeat_cost(ctx,monkeypatch):
    from Function.divine import divine
    from Function.divine.son_function.zhou_yi import from_lines
    monkeypatch.setattr(divine,"cast",lambda:from_lines([7]*6))
    # Even a negative legacy random draw must become a positive auspicious result.
    monkeypatch.setattr(divine,"draw_luck",lambda rng=None:-1)
    first=await divine.run(await task(ctx),ctx)
    assert "卦运 +1 · 福缘 1" in first["actions"][0]["text"]
    await ctx.db.write("members",[{"user_id":"123","fortune":2}],group_id=ctx.gid,mode="update",key_fields=("user_id",))
    t=await task(ctx)
    repeat=await divine.run(t,ctx)
    assert "卦运 +1 · 福缘 1" in repeat["actions"][0]["text"]
    assert await divine.run(t,ctx)==repeat
    record=(await ctx.db.read("divinations",group_id=ctx.gid,filters={"task_id":t["task_id"]}))[0]
    assert record["luck_delta"]==1
    assert record["zhou_yi"]["repeat_cost"]==2
    assert record["zhou_yi"]["luck_version"]=="readings-v1"
    assert record["zhou_yi"]["luck_judgment"]["direction"]==1
    assert record["tarot"]["orientation"]=="正位"
    await ctx.db.write("members",[{"user_id":"123","fortune":2}],group_id=ctx.gid,mode="update",key_fields=("user_id",))
    monkeypatch.setattr(divine,"cast",lambda:from_lines([7,7,7,7,7,9]))
    monkeypatch.setattr(divine,"draw_luck",lambda rng=None:3)
    bad=await divine.run(await task(ctx),ctx)
    assert "卦运 -3 · 福缘 -3" in bad["actions"][0]["text"]
